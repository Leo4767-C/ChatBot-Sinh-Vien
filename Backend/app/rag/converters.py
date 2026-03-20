from __future__ import annotations

import csv
import subprocess
# import tempfile
from pathlib import Path
import re, logging, datetime, os
from typing import Protocol

import os

import openpyxl
# import whisper
# from gtts import gTTS
# from marker.converters.pdf import PdfConverter
# from marker.models import create_model_dict
# from marker.output import text_from_rendered
# from pydub import AudioSegment

from app.core.config import settings

os.environ["CUDA_VISIBLE_DEVICES"] = settings.CUDA_VISIBLE_DEVICES

log = logging.getLogger(__name__)


class DocumentConverter(Protocol):
    def __call__(self, src: Path, dest: Path, *, to_console: bool = True) -> Path: ...


def _clean(text: str) -> str:
    text = re.sub(r"<br\s*/?>", " ", text)
    text = re.sub(r"<(.*?)>", r"[\1]", text)
    return text


# def pdf_to_markdown(src: Path, dest: Path, *, to_console: bool = True) -> Path:
#     """
#     Chuyển một file PDF sang markdown.
#
#     Parameters
#     ----------
#     src : Path
#         Đường dẫn tới PDF.
#     dest : Path
#         File .md muốn ghi.
#     to_console : bool, default=True
#         In log ra console hay không.
#
#     Returns
#     -------
#     Path
#         Đường dẫn file markdown đã sinh.
#     """
#     if not src.exists():
#         raise FileNotFoundError(src)
#
#     if to_console:
#         log.info("Start converting %s ➜ %s", src.name, dest.name)
#         log.info("Time: %s", datetime.datetime.now())
#
#     converter = PdfConverter(artifact_dict=create_model_dict())
#     rendered = converter(str(src))
#     text, _, _ = text_from_rendered(rendered)
#
#     text = _clean(text)
#
#     dest.write_text(text, encoding="utf-8")
#     log.info("✔ Markdown saved to %s", dest)
#     return dest

def convert_to_pdf(src: Path, dest: Path, *, to_console: bool = True) -> Path:
    """
    Convert document (docx/doc, pptx) to pdf
    Args:
        src: source path
        dest: destination path
        to_console: show log to console if True
    Returns: destination path
    """
    cmd = [
        settings.LIBRE_OFFICE,
        "--headless",
        "--convert-to", "pdf",
        str(src),
        "--outdir", str(dest.parent)
    ]

    if to_console:
        log.info("Start converting %s ➜ %s", src.name, dest.name)

    result = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    if result.returncode != 0:
        raise RuntimeError(f"LibreOffice convert error: {result.stderr.decode()}")

    log.info("✔ Pdf saved to %s", dest)
    return dest

def excel_to_csv(src: Path, dest_dir: Path, *, to_console: bool = True) -> list[Path]:
    """
    Convert Excel file to csv
    Args:
        src: Excel file path
        dest_dir: Destination directory to save csv files
        to_console: Show log to console if True
    Returns: List of csv paths
    """
    if to_console:
        log.info(f"Start converting {src} to csv and saving to {dest_dir}")

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    paths = []

    for sheet in wb.sheetnames:
        ws = wb[sheet]
        safe_sheet_name = sheet.replace("/", "_").replace("\\", "_")
        csv_path = dest_dir / f"{safe_sheet_name} - {src.stem}.csv"

        with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                writer.writerow([cell if cell is not None else "" for cell in row])
        paths.append(csv_path)

        if to_console:
            log.info(f"✔ Sheet '{sheet}' saved to {csv_path}")

    return paths


# def text_to_speech(text: str, lang: str = "vi", speed: float = 1.25):
#     try:
#         tts = gTTS(text=text, lang=lang, slow=False)
#         with tempfile.NamedTemporaryFile(suffix=".mp3", delete=False) as fp:
#             tts.save(fp.name)
#             # return fp.name
#
#             audio = AudioSegment.from_mp3(fp.name)
#             # Tăng tốc độ bằng cách thay đổi frame_rate
#             faster_audio = audio._spawn(audio.raw_data, overrides={
#                 "frame_rate": int(audio.frame_rate * speed)
#             })
#             faster_audio = faster_audio.set_frame_rate(audio.frame_rate) # Cập nhật lại frame_rate
#
#             with tempfile.NamedTemporaryFile(suffix=".mp3", delete=False) as tmp_faster_speed_file:
#                 faster_audio.export(tmp_faster_speed_file.name, format="mp3")
#                 return tmp_faster_speed_file.name
#     except Exception as e:
#         print(f"Error during gTTS conversion: {e}")
#         return None
#
#
# def speech_to_text(audio: str, lang: str = "vi"):
#     # DEVICE = "cuda" if torch.cuda.is_available() else "cpu"
#     model = whisper.load_model("base", device="cpu")
#     try:
#         audio = whisper.load_audio(audio)
#         audio = whisper.pad_or_trim(audio)
#         mel = whisper.log_mel_spectrogram(audio).to(model.device)
#
#         # Prepare decoding options
#         options = whisper.DecodingOptions(language=lang, fp16=True)
#
#         # Decode the audio
#         result = whisper.decode(model, mel, options)
#
#         # Return text
#         return result.text
#     except Exception as e:
#         print(e)
#         return None